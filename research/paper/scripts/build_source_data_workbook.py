#!/usr/bin/env python3
"""Build the consolidated Source Data workbook from archived result tables."""

from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEETS = [
    ("Fig2a_methods", "figure2/panel_a_method_comparison.tsv"),
    ("Fig2b_coverage", "figure2/panel_b_routing_metrics.tsv"),
    ("Fig2c_domains", "figure2/panel_c_domain_metrics.tsv"),
    ("Fig2_position", "figure2/position_accuracy_2530.tsv"),
    ("Fig2_cohort", "figure2/cohort_evaluation.tsv"),
    ("Fig3_ablation", "ablation_ladder_2530.tsv"),
    ("Fig3_cases", "agent_ablation_2530_case_changes.tsv"),
    ("Fig3_position_delta", "agent_ablation_2530_position_delta.tsv"),
    ("Fig4_AF3", "af3_cca13/af3_cca13_metrics.tsv"),
    ("Fig6_ENSURE", "ensure_perpos/perpos_sprinzlmap_edits.tsv"),
    ("Fig6_readthrough", "fig06_readthrough.tsv"),
    ("Fig6_latency", "fig06_latency.tsv"),
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def style_sheet(sheet) -> None:
    if sheet.max_row == 1:
        sheet.freeze_panes = "A2"
        return
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in column[: min(250, sheet.max_row)]]
        width = min(60, max(10, max((len(value) for value in values), default=10) + 2))
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--benchmark-root", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    source_rows = [
        {
            "Workbook sheet": "Cover",
            "Source file": "",
            "SHA-256": "",
            "Description": "Consolidated source-data workbook for SprinzlMap.",
        }
    ]
    for sheet_name, relative in SHEETS:
        path = args.root / relative
        if not path.exists():
            raise FileNotFoundError(path)
        source_rows.append(
            {
                "Workbook sheet": sheet_name,
                "Source file": str(path.relative_to(args.root)),
                "SHA-256": sha256(path),
                "Description": "Tabular source values retained verbatim from the archived result file.",
            }
        )
    if args.benchmark_root:
        for name, relative in [
            ("Gold provenance", "gold_2530/reference.provenance.json"),
            ("Run provenance", "runs/20261020-164533-deepseekv4-flash-guangdong/summary.json"),
            ("Run contract", "runs/20261020-164533-deepseekv4-flash-guangdong/manifest.json"),
        ]:
            path = args.benchmark_root / relative
            if path.exists():
                source_rows.append(
                    {
                        "Workbook sheet": name,
                        "Source file": str(path),
                        "SHA-256": sha256(path),
                        "Description": "Machine-readable provenance retained in the data archive.",
                    }
                )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        cover = pd.DataFrame(source_rows)
        cover.to_excel(writer, sheet_name="Cover", index=False)
        for sheet_name, relative in SHEETS:
            table = pd.read_csv(args.root / relative, sep="\t")
            table.to_excel(writer, sheet_name=sheet_name, index=False)
        for sheet in writer.book.worksheets:
            style_sheet(sheet)
    print(args.output)


if __name__ == "__main__":
    main()
